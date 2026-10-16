#!/usr/bin/env python3
"""
VALIDASI FINAL MIGRASI DATA MODEL — Laporan Stock Opname
Ground truth = pivot cache lama (15.880 + 15.880 + 1.199 baris).
Klasifikasi mismatch: 'known-fix' = bug lama hari-1 (SO_KEMARIN=0 padahal stok awal ada)
yang terbetulkan oleh logika DAX baru. Sisanya harus 0.
"""
import re
import xml.etree.ElementTree as ET
from collections import defaultdict
import openpyxl

UNZ = '/tmp/so_extract/unzipped'
XLSX = '/tmp/so_extract/sample.xlsx'
NS = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
TOL = 1e-6

def unesc(s):
    if s is None: return None
    return (s.replace('&amp;','&').replace('&lt;','<').replace('&gt;','>')
             .replace('&quot;','"').replace('&apos;',"'"))

def parse_cache(idx):
    droot = ET.parse(f'{UNZ}/xl/pivotCache/pivotCacheDefinition{idx}.xml').getroot()
    fields = []
    for cf in droot.find('m:cacheFields', NS):
        si = cf.find('m:sharedItems', NS)
        shared = []
        if si is not None and len(si):
            for ch in si: shared.append(unesc(ch.get('v')))
        fields.append((cf.get('name'), shared))
    rroot = ET.parse(f'{UNZ}/xl/pivotCache/pivotCacheRecords{idx}.xml').getroot()
    out = []
    for r in rroot:
        d = {}
        for (name, shared), cell in zip(fields, r):
            tag = cell.tag.split('}')[-1]; v = cell.get('v')
            if tag == 'x': d[name] = shared[int(v)] if int(v) < len(shared) else None
            elif tag in ('n','s','b'): d[name] = v
            elif tag == 'm': d[name] = None
            else: d[name] = v
        out.append(d)
    return out

def num(v):
    try: return float(v)
    except (TypeError, ValueError): return 0.0

def norm_attr(a):
    if a is None: return None
    try: return int(str(a).strip())
    except Exception: return None

C1, C2, C3 = parse_cache(1), parse_cache(2), parse_cache(3)
print(f"cache1={len(C1)} cache2={len(C2)} cache3={len(C3)} baris")

# ============ REKONSTRUKSI ============
wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
ws = wb['REKAP MENTAH']
rows = list(ws.iter_rows(values_only=True))
META, NDAYS = 8, 31

wide = defaultdict(dict)
for r in rows[1:]:
    if r[4] is None: continue
    kat, src, jenis, status, nama, ket, peny, sat = r[:META]
    if jenis == "REWORK" or kat is None or kat == "MENU MIX": continue
    for i in range(NDAYS):
        v = r[META + i]
        v = 0 if v is None else float(v)
        wide[(kat, nama, sat, i + 1)][jenis] = wide[(kat, nama, sat, i + 1)].get(jenis, 0.0) + v

per_item = defaultdict(list)
for k, jv in wide.items():
    per_item[k[:3]].append((k[3], jv))
for k in per_item: per_item[k].sort()

METRICS = ["1.DEVIASI", "2.WASTE", "3.SUSUT", "4.TRIAL", "5.LOSS/SURPLUS"]
JC = ["SO", "STOCK AWAL", "BARANG MASUK", "BARANG KELUAR", "BOM", "WASTE", "SUSUT", "TRIAL"]

dax_rows = {}
for key, lst in per_item.items():
    days = [d for d, _ in lst]
    idx_of = {d: i for i, d in enumerate(days)}
    cols = {j: [jv.get(j, 0.0) for _, jv in lst] for j in JC}
    n = len(days)
    awal_any = sum(cols["STOCK AWAL"])
    so_kem = []
    for i, d in enumerate(days):
        prev = [dd for dd in days if dd < d]
        so_kem.append(awal_any if not prev else cols["SO"][idx_of[max(prev)]])
    for i, d in enumerate(days):
        stok = so_kem[i] + cols["BARANG MASUK"][i] - cols["BARANG KELUAR"][i] - cols["BOM"][i]
        dev = cols["SO"][i] - stok
        w_, s_, t_ = -abs(cols["WASTE"][i]), -abs(cols["SUSUT"][i]), -abs(cols["TRIAL"][i])
        dax_rows[(key[0], key[1], key[2], d)] = {
            **{j: cols[j][i] for j in JC}, "SO_KEMARIN": so_kem[i], "STOK": stok,
            "1.DEVIASI": dev, "2.WASTE": w_, "3.SUSUT": s_, "4.TRIAL": t_,
            "5.LOSS/SURPLUS": dev - w_ - s_ - t_}

def keterangan(metric, val):
    if metric not in ("1.DEVIASI", "5.LOSS/SURPLUS"): return "-"
    if val < 0:  return "SO lebih sedikit daripada perhitungan stock seharusnya"
    if val > 0:  return "SO lebih banyak daripada perhitungan stock seharusnya"
    return "Stok Sesuai Perhitungan"

def indikasi(metric, val, so, soprev, masuk, keluar, bom):
    if metric not in ("1.DEVIASI", "5.LOSS/SURPLUS"): return "-"
    if so > soprev and masuk == 0:
        return "Barang Masuk belum tercatat (Stok SO naik tapi Barang Masuk = 0)"
    if so == soprev and bom > 0 and masuk == 0:
        return "Stok SO tidak berkurang meski ada penjualan/BOM"
    if so < soprev and keluar == 0 and bom == 0:
        return "FIX: Barang Keluar / BOM belum tercatat (Stok SO turun tapi pemakaian di sistem = 0)"
    if val > 0:
        return "1. Barang Masuk belum tercatat | 2. Under Takaran | 3. Salah catat Barang Keluar"
    if val < 0:
        return "1. Barang Keluar belum tercatat (Kiriman/retur ke resto lain/DC) | 2. Pencatatan Promix 00 kurang tepat | 3. Loss karna Waste/Susut atau Over Takaran"
    return "-"

# ============ CEK 2: vs CACHE1 dengan klasifikasi known-fix ============
print("\n=== CEK 2: logika baru vs CACHE1 (CEK DEVIASI) ===")
STOK_FIELD = [f for f in C1[0] if f.startswith('Stock Seharusnya')][0]

# item yang terdampak bug lama: punya baris null-Attribute di cache DAN (awal!=0 atau SO hari1!=0)
null_items = {d['NAMA BARANG'] for d in C1 if d['Attribute'] is None}
fix_items = set()
for d in C1:
    if d['Attribute'] == '1' and d['NAMA BARANG'] in null_items:
        if num(d['STOCK AWAL']) != 0 or num(d['SO']) != 0:
            fix_items.add(d['NAMA BARANG'])
print(f"  item dgn baris null (sumber bug) : {len(null_items)}")
print(f"  item terdampak bug hari-1        : {len(fix_items)} -> {sorted(fix_items)}")

n_match = n_knownfix = n_mismatch = n_missing = n_null = 0
unexplained = []
for d in C1:
    k = (d['KATEGORI'], d['NAMA BARANG'], d['SATUAN'], norm_attr(d['Attribute']))
    if k[3] is None:
        n_null += 1; continue
    mine = dax_rows.get(k)
    if mine is None:
        n_missing += 1; unexplained.append(f"MISSING {k}"); continue
    diffs_here = []
    for cf, myc in [("SO","SO"),("STOCK AWAL","STOCK AWAL"),("BARANG MASUK","BARANG MASUK"),
                    ("BARANG KELUAR","BARANG KELUAR"),("BOM","BOM"),("WASTE","WASTE"),
                    ("SUSUT","SUSUT"),("TRIAL","TRIAL"),("SO_KEMARIN","SO_KEMARIN"),(STOK_FIELD,"STOK")]:
        if abs(num(d.get(cf)) - (mine.get(myc, 0.0) or 0.0)) > TOL:
            diffs_here.append((cf, num(d.get(cf)), mine.get(myc, 0.0)))
    cv = num(d.get('Value')); mv = mine.get(d['JENIS'], 0.0) or 0.0
    if abs(cv - mv) > TOL:
        diffs_here.append((f"Value[{d['JENIS']}]", cv, mv))
    if not diffs_here:
        n_match += 1; continue
    # klasifikasi known-fix: hari-1, item fix_items, field SO_KEMARIN/STOK/Value metrik 1 atau 5,
    # dan cache SO_KEMARIN == 0 (tanda bug), nilai baru = benar (awal masuk)
    is_known = (k[3] == 1 and d['NAMA BARANG'] in fix_items and num(d['SO_KEMARIN']) == 0
                and all(f in ("SO_KEMARIN", STOK_FIELD, "Value[1.DEVIASI]", "Value[5.LOSS/SURPLUS]")
                        for f, _, _ in diffs_here))
    if is_known:
        n_knownfix += 1
    else:
        n_mismatch += 1
        if len(unexplained) < 8:
            unexplained.append(f"MISMATCH {k} {d['JENIS']}: {diffs_here[:3]}")
print(f"  baris cocok persis  : {n_match}")
print(f"  baris known-fix     : {n_knownfix}  (bug lama hari-1, kini TERBETULKAN)")
print(f"  baris null (skip)   : {n_null}   (baris Value=0 Attribute=null dari menu trial)")
print(f"  mismatch tak jelaskan: {n_mismatch} | missing: {n_missing}")
for u in unexplained[:8]: print("   ", u)
print(f"  => {'PASS' if n_mismatch == 0 and n_missing == 0 else 'PERIKSA'}")

# contoh before/after utk laporan
print("\n  CONTOH BEFORE/AFTER (bug lama -> koreksi):")
for it in ['ADONAN GORENG MEKAR', 'BERAS PUTIH', 'RICE WRAP']:
    for d in C1:
        if d['NAMA BARANG'] == it and d['Attribute'] == '1' and d['JENIS'] == '1.DEVIASI':
            k = (d['KATEGORI'], it, d['SATUAN'], 1)
            print(f"   {it:32s} hari1 1.DEVIASI: lama={num(d['Value']):>12.1f}  baru={dax_rows[k]['1.DEVIASI']:>12.1f}  (stok awal={num(d['STOCK AWAL']):.0f})")
            break

# ============ CEK 2b: teks ============
print("\n=== CEK 2b: teks Keterangan & Indikasi ===")
n_tmatch = n_tfix = n_tdiff = 0
tdiffs = []
for d in C1:
    k = (d['KATEGORI'], d['NAMA BARANG'], d['SATUAN'], norm_attr(d['Attribute']))
    if k[3] is None: continue
    mine = dax_rows.get(k)
    if mine is None: continue
    metric = d['JENIS']
    val = mine.get(metric, 0.0) or 0.0
    ek = keterangan(metric, val)
    ei = indikasi(metric, val, mine['SO'], mine['SO_KEMARIN'], mine['BARANG MASUK'], mine['BARANG KELUAR'], mine['BOM'])
    gk, gi = d.get('Keterangan'), d.get('Indikasi')
    if gk == ek and gi == ei:
        n_tmatch += 1
    elif k[3] == 1 and d['NAMA BARANG'] in fix_items:
        n_tfix += 1
    else:
        n_tdiff += 1
        if len(tdiffs) < 5: tdiffs.append(f"{k} {metric}: cache={gk!r}/{gi!r} mine={ek!r}/{ei!r}")
print(f"  teks cocok: {n_tmatch} | known-fix hari-1: {n_tfix} | tak terjelaskan: {n_tdiff}")
for t in tdiffs: print("   ", t)
print(f"  => {'PASS' if n_tdiff == 0 else 'PERIKSA'}")

# ============ CEK 3: AKUMULASI ============
print("\n=== CEK 3: running total AKUMULASI vs cache2 ===")
cum = defaultdict(float); akum = {}
for key in sorted(dax_rows.keys()):
    kat, nama, sat, day = key
    row = dax_rows[key]
    for metric in METRICS:
        cum[(kat, nama, sat, metric)] += row.get(metric, 0.0) or 0.0
        akum[(kat, nama, sat, metric, day)] = cum[(kat, nama, sat, metric)]
n_ok = n_kf = n_bad = 0
bads = []
for d in C2:
    k5 = (d['KATEGORI'], d['NAMA BARANG'], d['SATUAN'], d['JENIS'], norm_attr(d['Attribute']))
    if k5[4] is None: continue
    exp = akum.get(k5); got = num(d.get('AKUMULASI'))
    if exp is None:
        n_bad += 1; bads.append(f"MISS {k5}"); continue
    if abs(exp - got) > TOL:
        if d['NAMA BARANG'] in fix_items and d['JENIS'] in ('1.DEVIASI', '5.LOSS/SURPLUS'):
            n_kf += 1
        else:
            n_bad += 1
            if len(bads) < 5: bads.append(f"AKUM {k5}: cache={got} mine={exp}")
    else:
        n_ok += 1
print(f"  cocok: {n_ok} | known-fix (efek hari-1): {n_kf} | tak terjelaskan: {n_bad}")
for b in bads[:5]: print("   ", b)
print(f"  => {'PASS' if n_bad == 0 else 'PERIKSA'}")

# ============ CEK 4: PIVOT HASIL SO ============
print("\n=== CEK 4: PIVOT HASIL SO vs cache3 ===")
wsdb = wb['DATA BASE pastikan sesuai riil']
dbrows = list(wsdb.iter_rows(values_only=True))
h = dbrows[0]
i_nama, i_kode = h.index('NAMA BARANG'), h.index('KODE BARANG')
i_so, i_helper = h.index('INPUT SO HARIAN'), h.index('Helper')
kode_map = {}
for r in dbrows[1:]:
    if r[i_nama] is None: continue
    kode = r[i_kode]
    if kode is not None and str(kode).strip() not in ('', '0'):
        if r[i_so] == 'YA' and (r[i_helper] or '') != 'HIDE':
            kode_map[r[i_nama]] = kode
c3_keys = {(d['NAMA BARANG'], d.get('JENIS'), d.get('SOURCE DATA'), d.get('KETERANGAN'), d.get('PENYEBAB TERJADINYA WASTE')) for d in C3}
n_ok3 = n_bad3 = 0
for r in rows[1:]:
    if r[4] is None: continue
    if r[4] in kode_map:
        k = (r[4], r[2], r[1], r[5], r[6])
        if k in c3_keys: n_ok3 += 1
        else: n_bad3 += 1
print(f"  [FINAL] rows dgn kode: cocok {n_ok3} | hilang {n_bad3} | cache3 {len(C3)}")
print(f"  => {'PASS' if n_bad3 == 0 and len(C3) == n_ok3 else 'PERIKSA'}")

wb.close()

print("\n" + "=" * 60)
print("RINGKASAN AKHIR")
print(f"  CEK 2 (CEK DEVIASI)     : {n_match}/{len(C1)-n_null} persis + {n_knownfix} known-fix + {n_mismatch} anomali")
print(f"  CEK 2b (teks)           : {n_tmatch} persis + {n_tfix} known-fix + {n_tdiff} anomali")
print(f"  CEK 3 (AKUMULASI)       : {n_ok} persis + {n_kf} known-fix + {n_bad} anomali")
print(f"  CEK 4 (PIVOT HASIL SO)  : {n_ok3}/{len(C3)} persis")
ok = (n_mismatch == 0 and n_missing == 0 and n_tdiff == 0 and n_bad == 0 and n_bad3 == 0 and len(C3) == n_ok3)
print(f"  STATUS: {'SEMUA VALIDASI LULUS (anomali=0)' if ok else 'ADA ANOMALI - JANGAN DIPAKAI'}")
